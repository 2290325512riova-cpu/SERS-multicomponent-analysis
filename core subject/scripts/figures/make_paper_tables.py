from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

from paper_style import ROOT, TABLE_DIR, TASK_LABEL, TASK_ORDER


PAPER = ROOT / "data" / "pure63_mainline" / "models" / "paper_main"
SPLIT = ROOT / "data" / "pure63_mainline" / "splits" / "cv_split_random_5fold.csv"


def write_table(df: pd.DataFrame, stem: str) -> None:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    df.to_csv(TABLE_DIR / f"{stem}.csv", index=False)
    try:
        (TABLE_DIR / f"{stem}.md").write_text(df.to_markdown(index=False), encoding="utf-8")
    except Exception:
        (TABLE_DIR / f"{stem}.md").write_text(df.to_string(index=False), encoding="utf-8")
    (TABLE_DIR / f"{stem}.tex").write_text(df.to_latex(index=False, escape=False), encoding="utf-8")


def table_1_peak_assignment() -> pd.DataFrame:
    rows = [
        ("Thiram", 560, "S-S stretching / skeletal mode", "Published thiram SERS/Raman assignments", "core marker"),
        ("Thiram", 860, "(CH3)2-N / CH3 mode", "RSC Adv. 2024 thiram band near 881; normal Raman near 850", "assigned Thiram band"),
        ("Thiram", 930, "C-S symmetric stretching", "RSC Adv. 2024; Kang 2017", "core marker"),
        ("Thiram", 1145, "C-N stretching", "Published thiram SERS assignments", "core marker"),
        ("Thiram", 1380, "CH3 deformation / C-N", "Published thiram SERS assignments", "overlap-sensitive region"),
        ("Thiram", 1444, "C-N stretching + CH3 rocking", "Analyst 2014; RSC Adv. 2024 thiram band near 1436", "assigned Thiram band"),
        ("Thiram", 1510, "C=N / C-N related mode", "Kang 2017; Zhu 2019", "core marker"),
        ("MG", 908, "ring vibration", "Published MG SERS assignments", "marker"),
        ("MG", 1172, "C-H in-plane bending", "Published MG SERS assignments", "MG-associated overlap-sensitive region"),
        ("MG", 1220, "C-H in-plane bending", "RSC Adv. 2014; Jiang 2018", "marker"),
        ("MG", 1394, "aromatic C-C / C-N mode", "Published MG SERS assignments", "MG-associated overlap-sensitive region"),
        ("MG", 1616, "aromatic C=C stretching", "Published MG SERS assignments; this work Fig.5", "clean MG suppression marker"),
        ("4-MBA", 1078, "ring breathing / C-S related mode", "Published 4-MBA SERS assignments", "core marker"),
        ("4-MBA", 1180, "C-H bending", "Published 4-MBA SERS assignments", "overlap-sensitive region"),
        ("4-MBA", 1490, "aromatic C-C stretching", "Published 4-MBA SERS assignments", "marker"),
        ("4-MBA", 1590, "aromatic C=C stretching", "Published 4-MBA SERS assignments", "core marker"),
    ]
    return pd.DataFrame(
        rows,
        columns=["Analyte", "Raman shift (cm^-1)", "Assignment", "Literature / evidence support", "Manuscript role"],
    )


def table_2_best_by_task() -> pd.DataFrame:
    df = pd.read_csv(PAPER / "benchmark_random_cv" / "benchmark_best_by_task.csv")
    df = df.set_index("Task").reindex(TASK_ORDER).reset_index()
    out = pd.DataFrame(
        {
            "Task": [TASK_LABEL[t] for t in df.Task],
            "Best model": df.Model,
            "Preprocess": df.Preprocess,
            "Macro-F1": df.apply(lambda r: f"{r.F1_mean:.3f} +/- {r.F1_std:.3f}", axis=1),
            "Balanced accuracy": df.apply(lambda r: f"{r.BA_mean:.3f} +/- {r.BA_std:.3f}", axis=1),
            "Accuracy": df.apply(lambda r: f"{r.Acc_mean:.3f} +/- {r.Acc_std:.3f}", axis=1),
        }
    )
    return out


def table_3_soil_metrics() -> pd.DataFrame:
    summary = pd.read_csv(PAPER / "soil_validation" / "soil_cv_summary.csv")
    blank = pd.read_csv(PAPER / "soil_validation" / "soil_blank_specificity.csv")
    df = summary.merge(blank[["task", "specificity"]], on="task")
    df = df.set_index("task").reindex(["P1_thiram_presence", "P2_mg_presence", "P3_mba_presence"]).reset_index()
    return pd.DataFrame(
        {
            "Task": [TASK_LABEL[t] for t in df.task],
            "Soil spectra": df.n_test_total,
            "Blank controls": df.n_blank_test_total,
            "AUC": df.apply(lambda r: f"{r.AUC_mean:.3f} +/- {r.AUC_std:.3f}", axis=1),
            "Macro-F1": df.apply(lambda r: f"{r.F1_mean:.3f} +/- {r.F1_std:.3f}", axis=1),
            "Balanced accuracy": df.apply(lambda r: f"{r.balanced_acc_mean:.3f} +/- {r.balanced_acc_std:.3f}", axis=1),
            "Blank specificity": df.apply(lambda r: f"{r.specificity:.3f}", axis=1),
        }
    )


def table_representative_rsd() -> pd.DataFrame:
    meta = pd.read_csv(SPLIT)
    processed = ROOT / "data" / "pure63_mainline" / "processed"
    wn = np.load(processed / "wavenumber.npy")
    x = np.load(processed / "X_p1.npy")
    specs = [
        ("Thiram", "Thiram_1382", 1382, "has_thiram", 4),
        ("MG", "MG_1616", 1616, "has_mg", 4),
        ("4-MBA", "MBA_1080", 1080, "has_mba", 4),
    ]
    rows = []
    for analyte, peak_name, peak_wn, flag, level in specs:
        if analyte == "Thiram":
            mask = (meta.has_thiram == 1) & (meta.has_mg == 0) & (meta.has_mba == 0) & (meta.c_thiram == level)
        elif analyte == "MG":
            mask = (meta.has_mg == 1) & (meta.has_thiram == 0) & (meta.has_mba == 0) & (meta.c_mg == level)
        else:
            mask = (meta.has_mba == 1) & (meta.has_thiram == 0) & (meta.has_mg == 0) & (meta.c_mba == level)
        idx = np.where(mask.to_numpy())[0]
        band = (wn >= peak_wn - 8) & (wn <= peak_wn + 8)
        areas = np.trapezoid(x[idx][:, band], wn[band], axis=1)
        mean_area = float(np.mean(areas))
        rsd_percent = float(np.std(areas, ddof=1) / abs(mean_area) * 100)
        rows.append(
            {
                "Analyte": analyte,
                "Marker peak (cm^-1)": peak_wn,
                "Concentration": f"1e-{level} M",
                "Preprocess": "p1",
                "n spectra": len(idx),
                "Metric": "marker-band area (+/-8 cm^-1)",
                "Mean band area": f"{mean_area:.4g}",
                "RSD (%)": f"{rsd_percent:.2f}",
                "Figure role": "Fig.2b same-concentration pure-standard repeatability",
            }
        )
    return pd.DataFrame(rows)


def supplementary_tables() -> dict[str, pd.DataFrame]:
    tables: dict[str, pd.DataFrame] = {}
    audit = pd.read_csv(PAPER / "shap_peak_cluster" / "shap_assignment_audit.csv")
    write_table(audit, "ST1_shap_assignment_audit")
    tables["ST1_SHAP_audit"] = audit

    full = pd.read_csv(PAPER / "benchmark_random_cv" / "benchmark_full_matrix.csv")
    write_table(full, "ST2_full_benchmark_matrix")
    tables["ST2_full_benchmark"] = full

    fs = pd.read_csv(PAPER / "shap_feature_selection" / "shap_feature_elimination_curve.csv")
    write_table(fs, "ST3_shap_feature_selection_curve")
    tables["ST3_feature_selection"] = fs

    perm = pd.read_csv(PAPER / "soil_validation" / "soil_permutation_test.csv")
    write_table(perm, "ST4_soil_permutation_summary")
    tables["ST4_soil_permutation"] = perm

    rsd = table_representative_rsd()
    write_table(rsd, "ST5_representative_peak_rsd")
    tables["ST5_representative_RSD"] = rsd
    return tables


def write_excel_workbook(tables: dict[str, pd.DataFrame]) -> None:
    path = TABLE_DIR / "paper_tables_compiled.xlsx"
    with pd.ExcelWriter(path, engine="xlsxwriter") as writer:
        for sheet, df in tables.items():
            safe_sheet = sheet[:31]
            df.to_excel(writer, sheet_name=safe_sheet, index=False)
            workbook = writer.book
            worksheet = writer.sheets[safe_sheet]
            header_fmt = workbook.add_format({"bold": True, "bg_color": "#D9EAF7", "border": 1})
            for col_num, value in enumerate(df.columns):
                worksheet.write(0, col_num, value, header_fmt)
                width = min(max(12, int(df[value].astype(str).str.len().quantile(0.90)) + 2), 42)
                worksheet.set_column(col_num, col_num, width)
            worksheet.freeze_panes(1, 0)


def write_main_tables_docx(tables: dict[str, pd.DataFrame]) -> None:
    try:
        from docx import Document
        from docx.shared import Pt
    except Exception:
        return
    doc = Document()
    doc.add_heading("Main manuscript tables", level=1)
    captions = {
        "T1_peak_assignment": "Table 1. Characteristic SERS bands and their manuscript roles.",
        "T2_best_by_task": "Table 2. Best model and preprocessing combination for each screening task.",
        "T3_soil_metrics": "Table 3. Spiked-soil matrix screening performance for pesticide presence tasks.",
    }
    for key in ["T1_peak_assignment", "T2_best_by_task", "T3_soil_metrics"]:
        df = tables[key]
        doc.add_paragraph(captions[key])
        table = doc.add_table(rows=1, cols=len(df.columns))
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, col in enumerate(df.columns):
            hdr[i].text = str(col)
        for _, row in df.iterrows():
            cells = table.add_row().cells
            for i, col in enumerate(df.columns):
                cells[i].text = str(row[col])
        doc.add_paragraph("")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9)
    doc.save(TABLE_DIR / "main_tables_for_manuscript.docx")


def main() -> None:
    main_tables = {
        "T1_peak_assignment": table_1_peak_assignment(),
        "T2_best_by_task": table_2_best_by_task(),
        "T3_soil_metrics": table_3_soil_metrics(),
    }
    write_table(main_tables["T1_peak_assignment"], "T1_peak_assignment")
    write_table(main_tables["T2_best_by_task"], "T2_best_by_task_performance")
    write_table(main_tables["T3_soil_metrics"], "T3_soil_screening_metrics")
    supp = supplementary_tables()
    write_excel_workbook({**main_tables, **supp})
    write_main_tables_docx(main_tables)
    print(f"Wrote tables to {TABLE_DIR}")


if __name__ == "__main__":
    main()
